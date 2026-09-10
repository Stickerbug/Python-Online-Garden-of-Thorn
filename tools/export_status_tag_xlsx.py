# -*- coding: utf-8 -*-
"""导出“当前全部状态+标签”到单个 xlsx（PvP/多人，不含故事模式）。

内置定义来自 static/js/game.js（经 Node 沙箱执行提取），模组自定义来自
mods/*.gtnmod；来源标注为原版或所属模组中文名。

运行：python tools/export_status_tag_xlsx.py   （需要本机 node 可用）
"""
import io
import json
import os
import subprocess
import sys
import tempfile

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, ROOT)

JS_EXTRACT = r"""
const fs = require("fs");
const vm = require("vm");
const src = fs.readFileSync(process.argv[2], "utf8");
const noop = () => {};
const storeMap = {};
const storage = { getItem: k => (k in storeMap ? storeMap[k] : null), setItem: (k, v) => { storeMap[k] = String(v); }, removeItem: k => { delete storeMap[k]; } };
const classList = { add: noop, remove: noop, toggle: noop, contains: () => false };
const elem = new Proxy({}, { get: (t, k) => {
  if (k === "style") return {};
  if (k === "classList") return classList;
  if (k === "dataset") return {};
  if (k === "addEventListener" || k === "removeEventListener") return noop;
  if (k === "querySelectorAll") return [];
  if (k === "querySelector") return null;
  if (k === "appendChild") return elem;
  return elem;
}, set: () => true });
const win = { addEventListener: noop, removeEventListener: noop, location: { href: "http://x/", search: "", protocol: "http:" } };
const doc = { getElementById: () => null, querySelector: () => null, querySelectorAll: () => [], createElement: () => elem, addEventListener: noop, removeEventListener: noop, body: elem, documentElement: { classList, style: {}, addEventListener: noop }, hidden: false };
const sandbox = {
  console, document: doc, window: win, navigator: { userAgent: "node" },
  localStorage: storage, sessionStorage: storage, location: win.location, history: {},
  requestAnimationFrame: () => 0, cancelAnimationFrame: noop, setTimeout, clearTimeout,
  setInterval: () => 0, clearInterval, performance: { now: () => Date.now() },
  fetch: async () => ({ json: async () => ({}) }), addEventListener: noop, removeEventListener: noop,
  dispatchEvent: () => true, WebSocket: function () {}, Event: function () {}, CustomEvent: function () {},
  Image: function () {}, Audio: function () { return { play: noop, pause: noop }; },
  URL, URLSearchParams, Blob: function () {}, FormData: function () {}, FileReader: function () {},
  MutationObserver: function () { return { observe: noop, disconnect: noop }; },
};
sandbox.globalThis = sandbox;
vm.createContext(sandbox);
vm.runInContext(src, sandbox, { timeout: 60000 });
const registry = JSON.parse(fs.readFileSync(process.argv[3], "utf8"));
const customTags = {};
const customStatuses = {};
for (const mod of registry) {
  for (const tag of mod.tags || []) if (tag && tag.id) customTags[tag.id] = tag;
  for (const st of mod.statuses || []) if (st && st.id) customStatuses[st.id] = st;
}
vm.runInContext("setCustomRegistries(" + JSON.stringify(Object.values(customTags)) + "," + JSON.stringify(Object.values(customStatuses)) + ")", sandbox);
const out = {};
out.statuses = [...vm.runInContext("getAllStatusDefs()", sandbox).values()].map(s => ({ key: s.key, label: s.label || "", source: s.source || "", desc: s.desc || "" }));
out.flags = vm.runInContext("getAllGalleryFlags()", sandbox);
out.flagInfos = {};
for (const flag of out.flags) {
  try {
    const custom = vm.runInContext(`getCustomTagDef(${JSON.stringify(flag)})`, sandbox);
    out.flagInfos[flag] = {
      label: vm.runInContext(`getFlagLabel(${JSON.stringify(flag)})`, sandbox),
      desc: vm.runInContext(`getIntroFlagDescription(${JSON.stringify(flag)}, ${JSON.stringify(custom)})`, sandbox),
      customId: custom ? custom.id : null,
    };
  } catch (e) {
    out.flagInfos[flag] = { label: "", desc: "", customId: null };
  }
}
fs.writeFileSync(process.argv[4], JSON.stringify(out), "utf8");
"""


def main():
    from mod_loader import load_all_mods
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter

    game_js = os.path.join(ROOT, "static", "js", "game.js")
    out_dir = os.path.dirname(ROOT)
    out_xlsx = os.path.join(out_dir, "模组状态标签数据.xlsx")

    registry = []
    for mod in load_all_mods():
        source = {
            "mod_id": getattr(mod.manifest, "id", "") if mod.manifest else "",
            "filename": mod.filename,
            "name_cn": (mod.info.name_cn if mod.info else "") or "",
            "name_en": (mod.info.name_en if mod.info else mod.info.name if mod.info else "") or "",
        }
        tags = [r.to_dict() for r in (mod.registries.get("tags") or [])]
        statuses = [r.to_dict() for r in (mod.registries.get("statuses") or [])]
        if tags or statuses:
            registry.append({**source, "tags": tags, "statuses": statuses})

    tmp = tempfile.mkdtemp(prefix="gtn_export_")
    reg_path = os.path.join(tmp, "registry.json")
    js_path = os.path.join(tmp, "extract.js")
    dump_path = os.path.join(tmp, "dump.json")
    try:
        with open(reg_path, "w", encoding="utf-8") as f:
            json.dump(registry, f, ensure_ascii=False)
        with open(js_path, "w", encoding="utf-8") as f:
            f.write(JS_EXTRACT)
        proc = subprocess.run(
            ["node", js_path, game_js, reg_path, dump_path],
            check=True,
            capture_output=True,
            text=True,
        )
        dump = json.load(open(dump_path, encoding="utf-8"))
    finally:
        for path in (reg_path, js_path, dump_path):
            try:
                os.remove(path)
            except OSError:
                pass
        try:
            os.rmdir(tmp)
        except OSError:
            pass

    id_mod = {}
    for mod in registry:
        name = mod.get("name_cn") or mod.get("filename") or ""
        for tag in mod.get("tags") or []:
            if tag.get("id"):
                id_mod[tag["id"]] = name
        for st in mod.get("statuses") or []:
            if st.get("id"):
                id_mod[st["id"]] = name

    def resolve_source(key, front_source):
        if id_mod.get(key):
            return id_mod[key]
        return "原版"

    wb = Workbook()
    ws = wb.active
    ws.title = "状态与标签"
    ws.append(["来源", "类别", "ID", "中文名", "描述"])
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="DDEBF7")
    ws.append(["状态", "", "", "", ""])
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=5)
    ws.cell(2, 1).font = Font(bold=True)
    ws.cell(2, 1).fill = PatternFill("solid", fgColor="E2EFDA")
    for s in dump["statuses"]:
        ws.append([resolve_source(s["key"], s.get("source", "")), "状态", s["key"], s.get("label") or "", s.get("desc") or ""])
    ws.append([])
    tag_row = ws.max_row + 1
    ws.append(["标签", "", "", "", ""])
    ws.merge_cells(start_row=tag_row, start_column=1, end_row=tag_row, end_column=5)
    ws.cell(tag_row, 1).font = Font(bold=True)
    ws.cell(tag_row, 1).fill = PatternFill("solid", fgColor="FFF2CC")
    for flag in dump["flags"]:
        info = dump["flagInfos"].get(flag, {})
        src = "mod" if info.get("customId") else "vanilla"
        ws.append([resolve_source(flag, src), "标签", flag, info.get("label") or "", info.get("desc") or ""])
    for idx, width in enumerate([22, 8, 34, 22, 90], start=1):
        ws.column_dimensions[get_column_letter(idx)].width = width
    wb.save(out_xlsx)
    print("saved:", out_xlsx)
    print("statuses:", len(dump["statuses"]), "tags:", len(dump["flags"]))


if __name__ == "__main__":
    main()
