# -*- coding: utf-8 -*-
"""导出全部模组卡牌到单个 xlsx（与游戏图鉴一致的顺序）。

列：所属模组 / 名称(英) / 中文名称 / 消耗E / 消耗M / 类型(TBRG) /
标签（中文，逗号分隔） / 效果描述 / 中文描述

运行：python tools/export_cards_xlsx.py
"""
import os
import sys

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, ROOT)

from mod_loader import load_all_mods
from cards import normalize_card_flag

OUT = os.path.join(os.path.dirname(ROOT), "模组卡牌数据.xlsx")
TYPE_ORDER = {"thorn": 0, "bloom": 1, "guard": 2, "root": 3}
TYPE_LETTER = {"thorn": "T", "bloom": "B", "guard": "G", "root": "R"}

BUILTIN_TAG_CN = {
    "precision": "精准", "exile": "放逐", "non_stackable": "不叠加",
    "indestructible": "不可摧毁", "sprout": "萌芽", "symbiosis": "共生",
    "attract": "吸附", "void": "虚无", "self_only": "不选择目标",
    "uncancellable": "不可取消", "infinite_exclude": "无限火力移除",
    "rebound": "回转", "copy": "副本", "unique": "唯一", "swift": "迅捷",
    "temp_swift": "暂时迅捷", "temp_heavy": "暂时沉重",
    "temp_magic_heavy": "暂时魔力沉重", "floating": "漂浮",
    "stealth": "隐匿", "revealed": "被揭示", "sublime": "崇高",
    "team_limited": "队伍限定", "team_unique": "队伍独一", "power": "威力",
    "magic_swift": "魔力迅捷", "wide_strike": "广域打击", "self_target": "自刃",
    "charge": "电荷", "ocean_blinded": "蒙蔽", "amplify": "增幅",
    "fusion_layer": "聚变", "fission_layer": "裂变",
}


def load_custom_tag_names():
    names = {}
    for mod in load_all_mods():
        for res in mod.registries.get("tags") or []:
            data = res.to_dict()
            raw = data.get("id") or data.get("key") or ""
            norm = normalize_card_flag(raw)
            if not norm:
                continue
            names.setdefault(norm, str(
                data.get("name_i18n", {}).get("zh")
                or data.get("name_cn") or data.get("name") or data.get("name_en") or raw
            ))
    return names


def main():
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter

    custom = load_custom_tag_names()
    mods = load_all_mods()
    wb = Workbook()
    ws = wb.active
    ws.title = "模组卡牌数据"
    header = ["所属模组", "名称", "中文名称", "消耗E", "消耗M", "类型", "标签", "效果描述", "中文描述"]
    ws.append(header)
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="DDEBF7")
    first_mod = True
    total = 0
    for mod in mods:
        cards = sorted(
            (c.to_dict() for c in mod.cards),
            key=lambda c: (TYPE_ORDER.get(str(c.get("card_type") or ""), 99), str(c.get("id") or "").lower()),
        )
        if not cards:
            continue
        if not first_mod:
            ws.append([])
        first_mod = False
        mod_cn = (mod.info.name_cn if mod.info else "") or mod.filename
        for c in cards:
            labels = []
            for raw in c.get("flags") or []:
                norm = normalize_card_flag(raw)
                if not norm:
                    continue
                name = BUILTIN_TAG_CN.get(norm) or custom.get(norm) or norm
                if name not in labels:
                    labels.append(name)
            ws.append([
                mod_cn,
                c.get("name_en") or c.get("id") or "",
                c.get("name_cn") or "",
                c.get("cost_e"), c.get("cost_m"),
                TYPE_LETTER.get(str(c.get("card_type") or ""), str(c.get("card_type") or "")),
                "，".join(labels),
                c.get("effect_text") or "",
                c.get("description") or "",
            ])
            total += 1
    for idx, width in enumerate([24, 28, 20, 10, 10, 10, 34, 70, 70], start=1):
        ws.column_dimensions[get_column_letter(idx)].width = width
    wb.save(OUT)
    print("saved:", OUT)
    print("mods:", len(mods), "card rows:", total)


if __name__ == "__main__":
    main()
